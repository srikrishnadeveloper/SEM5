import argparse
import json
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor


def _resolve(base_dir, path):
    if os.path.isabs(path):
        return path
    return os.path.join(base_dir, path)


def set_cell_shading(cell, fill):
    tcPr = cell._element.get_or_add_tcPr()
    tcPr.append(parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), fill)))


def add_code_table(doc, text, plain=False):
    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    cell = table.cell(0, 0)

    if plain:
        bg, fg = 'FFFFFF', '000000'
    else:
        bg, fg = '1E1E1E', 'D4D4D4'
    set_cell_shading(cell, bg)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15

    r, g, b = int(fg[0:2], 16), int(fg[2:4], 16), int(fg[4:6], 16)
    for line in text.rstrip('\n').splitlines():
        run = p.add_run(line + '\n')
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(r, g, b)


def add_output_table(doc, text, plain=False):
    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    cell = table.cell(0, 0)

    if plain:
        bg, fg = 'F7F9FC', '000000'
    else:
        bg, fg = 'F7F9FC', '000000'
    set_cell_shading(cell, bg)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15

    for line in text.rstrip('\n').splitlines():
        run = p.add_run(line + '\n')
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)


def add_data_table(doc, spec_table):
    headers = spec_table.get('headers', [])
    rows = spec_table.get('rows', [])
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(h)
        set_cell_shading(cell, 'DDE7F5')
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True

    for j, row in enumerate(rows):
        for i, v in enumerate(row):
            table.rows[j + 1].cells[i].text = str(v)

    for i, w in enumerate(spec_table.get('widths', [])):
        try:
            table.columns[i].width = Pt(w)
        except Exception:
            pass


def build_docx(spec, out_path):
    doc = Document()
    plain = spec.get('plain', False)

    # title page / header
    if spec.get('title'):
        t = doc.add_heading(spec['title'], level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if spec.get('subtitle'):
        p = doc.add_paragraph(spec['subtitle'])
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.runs:
            p.runs[0].italic = True

    if any(k in spec for k in ('name', 'class_section', 'regno', 'institution')):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        parts = [
            spec.get('name', ''),
            spec.get('class_section', ''),
            spec.get('regno', ''),
            spec.get('institution', ''),
        ]
        p.add_run('   '.join([x for x in parts if x])).bold = True

    doc.add_paragraph()

    base_dir = os.path.dirname(os.path.abspath(out_path))

    for sec in spec.get('sections', []):
        if sec.get('page_break'):
            doc.add_page_break()
            continue

        if sec.get('heading'):
            h = doc.add_heading(sec['heading'], level=2)
        if sec.get('subheading'):
            doc.add_heading(sec['subheading'], level=3)

        if sec.get('paragraph'):
            paras = sec['paragraph'] if isinstance(sec['paragraph'], list) else [sec['paragraph']]
            for para in paras:
                doc.add_paragraph(para)

        if sec.get('code_file'):
            with open(_resolve(base_dir, sec['code_file']), encoding='utf-8') as f:
                add_code_table(doc, f.read(), plain=plain)
        elif sec.get('code'):
            add_code_table(doc, sec['code'], plain=plain)

        if sec.get('output_file'):
            with open(_resolve(base_dir, sec['output_file']), encoding='utf-8') as f:
                add_output_table(doc, f.read(), plain=plain)
        elif sec.get('output'):
            add_output_table(doc, sec['output'], plain=plain)

        if sec.get('table'):
            add_data_table(doc, sec['table'])

        images = sec.get('plots') or sec.get('screenshots')
        if images:
            for pl in images:
                img_path = pl.get('image', pl if isinstance(pl, str) else '')
                img_path = _resolve(base_dir, img_path)
                if os.path.isfile(img_path):
                    try:
                        doc.add_picture(img_path, width=Inches(6.15))
                    except Exception as e:
                        print('WARNING: could not add picture', img_path, e)
                        continue
                    if pl.get('caption'):
                        cp = doc.add_paragraph(pl['caption'])
                        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        if cp.runs:
                            cp.runs[0].italic = True

    if sec.get('bullets'):
        for b in sec['bullets']:
            doc.add_paragraph(b, style='List Bullet')

    doc.save(out_path)
    print('DOCX written:', out_path)


def main():
    parser = argparse.ArgumentParser(description='Build a lab-report DOCX from a spec.json')
    parser.add_argument('spec', help='path to spec.json')
    parser.add_argument('--out', default='report.docx', help='output DOCX path')
    args = parser.parse_args()

    if not os.path.isfile(args.spec):
        print('ERROR: spec file not found:', args.spec)
        sys.exit(1)

    with open(args.spec, encoding='utf-8') as f:
        spec = json.load(f)

    build_docx(spec, args.out)


if __name__ == '__main__':
    main()
