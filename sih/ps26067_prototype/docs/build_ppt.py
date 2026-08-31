"""Build an updated SIH26067 pitch deck."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BG = RGBColor(15, 23, 42)
ACCENT = RGBColor(56, 189, 248)
TEXT = RGBColor(226, 232, 240)
SUB = RGBColor(148, 163, 184)


def add_title_slide(prs, title, subtitle):
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)

    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()

    title_box = s.shapes.add_textbox(Inches(0.7), Inches(2.2), Inches(8.6), Inches(1.5))
    tf = title_box.text_frame
    tf.text = title
    tf.paragraphs[0].font.size = Pt(44)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = TEXT
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    sub_box = s.shapes.add_textbox(Inches(0.7), Inches(3.7), Inches(8.6), Inches(1.5))
    tf = sub_box.text_frame
    tf.text = subtitle
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.color.rgb = ACCENT
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    foot = s.shapes.add_textbox(Inches(0.7), Inches(6.8), Inches(8.6), Inches(0.6))
    tf = foot.text_frame
    tf.text = "Smart India Hackathon 2026 | Team [Name] | SSN College of Engineering"
    tf.paragraphs[0].font.size = Pt(12)
    tf.paragraphs[0].font.color.rgb = SUB
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER


def add_section_slide(prs, title, bullets):
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)

    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()

    title_box = s.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(8.8), Inches(0.9))
    tf = title_box.text_frame
    tf.text = title
    tf.paragraphs[0].font.size = Pt(32)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ACCENT

    content_box = s.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.4), Inches(5.0))
    tf = content_box.text_frame
    tf.word_wrap = True

    for i, b in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(20)
        p.font.color.rgb = TEXT
        p.space_after = Pt(14)


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "SIH26067 — OceanViz 3D",
        "A national-level web platform integrating numerical ocean model outputs and in-situ observations",
    )

    add_section_slide(
        prs,
        "Problem & Gaps",
        [
            "Ocean models and observations live in separate tools (2D maps, spreadsheets, desktop software).",
            "No single web-native 3D globe to co-visualise forecast fields with Argo / glider / CTD data.",
            "Missing depth slicing, time animation, current vectors and instrument profiles in one view.",
            "Result: slower analysis and limited situational awareness for fisheries, climate and marine ops.",
        ],
    )

    add_section_slide(
        prs,
        "Solution Architecture",
        [
            "Cesium.js 3D globe with free OpenStreetMap basemap and real lat/lon/depth mapping.",
            "FastAPI backend serves /image, /slice, /vectors, /instruments and /profile endpoints.",
            "Pluggable DataSource interface: swap SyntheticSource with NetCDF/OPeNDAP in one file.",
            "2026 glassmorphism UI language: frosted panels, depth, glow and microinteractions.",
            "Cesium & Chart.js vendored under static/lib/ for fully offline grand-finale demos.",
        ],
    )

    add_section_slide(
        prs,
        "Key Features",
        [
            "3D Cesium globe: pan, zoom, rotate over the Indian EEZ (72°–86°E, 8°–20°N).",
            "Ocean variables: temperature, salinity, chlorophyll, current speed with scientific colormaps.",
            "Depth slider: 0, 50, 100, 200, 500 m slices through the water column.",
            "Current vectors: U/V arrows coloured by speed, inspired by monsoon circulation.",
            "Time animation: play/pause through 6 synthetic forecast / observation steps.",
            "Click any instrument for a depth-vs-variable profile; click data layer to probe exact values.",
        ],
    )

    add_section_slide(
        prs,
        "Tech Stack",
        [
            "Frontend: vanilla JS, Cesium.js, Chart.js (all vendored for offline use).",
            "Backend: FastAPI, Python, NumPy, Pillow, Matplotlib.",
            "Data: 4D synthetic grid (lat/lon/depth/time) + instrument trajectories and profiles.",
            "Packaging: one-command run.ps1 / run.sh, Dockerfile and docker-compose.yml.",
            "Presentation: this PPT, demo script, architecture diagram and README.",
        ],
    )

    add_section_slide(
        prs,
        "Impact & Future",
        [
            "Impact: better fisheries planning, climate monitoring, marine operations and disaster response.",
            "Feasibility: open-source stack, no paid API keys, offline-first design.",
            "Modular DataSource means real INCOIS data is a one-file swap away.",
            "Future: cloud deployment, real-time OPeNDAP/THREDDS, multi-source fusion, ML insights.",
        ],
    )

    add_title_slide(
        prs,
        "Thank You",
        "OceanViz 3D — ready for the national grand finale",
    )

    out = __import__("pathlib").Path(__file__).with_name("SIH26067_Presentation.pptx")
    prs.save(out)
    print(f"Saved {out}")


if __name__ == "__main__":
    main()
